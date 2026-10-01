"""Planilha de quantitativo (seção 10.4). Custos por fórmula a partir da aba Precos."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..model import Result

F = "Arial"
HDR = PatternFill("solid", fgColor="1F1F1F")
INPUT = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BBBBBB")


def _hdr(ws, row, cols):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _font_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None,
                              size=c.font.size if c.font else 10)


def write_xlsx(res: Result, q: dict, cat, path: str) -> None:
    wb = Workbook()
    # ------------------------------------------------ Precos (entrada)
    wp = wb.active
    wp.title = "Precos"
    wp["A1"] = "Preços unitários (preencher células amarelas; valor vazio = 0)"
    wp["A1"].font = Font(name=F, bold=True, size=12)
    wp["A2"] = "Linear: R$ por metro de barra comprada · Placas: R$ por placa · Ferragens: R$ por unidade. Fonte: preencher com cotação do fornecedor e data."
    _hdr(wp, 4, ["Item", "Unidade de preço", "Preço (R$)", "Fornecedor", "Data da cotação"])
    items = [(r["item"], "m") for r in q["consolidated"]] + [(r["item"], "placa") for r in q["sheets"]] + \
            [(r["item"], "un") for r in q["hardware"]]
    seen = set()
    row = 5
    for it, un in items:
        if it in seen:
            continue
        seen.add(it)
        wp.cell(row=row, column=1, value=it)
        wp.cell(row=row, column=2, value=un)
        c = wp.cell(row=row, column=3, value=None)
        c.fill = INPUT
        c.font = Font(name=F, color="0000FF")
        c.number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
        for col in (4, 5):
            wp.cell(row=row, column=col).fill = INPUT
        row += 1
    last_price = row - 1
    wp.cell(row=5, column=3).comment = Comment("Exemplo de formato: 12,50. Preço de referência a confirmar.", "VFE")
    _widths(wp, [30, 16, 14, 22, 16])
    PR = f"Precos!$A$5:$A${last_price}"
    PV = f"Precos!$C$5:$C${last_price}"

    def price(ref):
        return f"IFERROR(INDEX({PV},MATCH({ref},{PR},0)),0)"

    # ------------------------------------------------ Materiais lineares
    wm = wb.create_sheet("Materiais")
    _hdr(wm, 1, ["Item", "Categoria", "Peças", "Metros de corte", "Barras (comprimento: qtd)", "Total de barras",
                 "Metros de barra", "Aproveitamento", "Sobras reaproveitáveis", "Peso (kg)", "Preço (R$/m)",
                 "Custo (R$)", "Barras mínimas teóricas"])
    for i, r in enumerate(q["consolidated"], start=2):
        wm.cell(row=i, column=1, value=r["item"])
        wm.cell(row=i, column=2, value=r["category"])
        wm.cell(row=i, column=3, value=r["pieces"])
        wm.cell(row=i, column=4, value=r["cut_m"]).number_format = "#,##0.00"
        wm.cell(row=i, column=5, value="; ".join(f"{int(k)}: {v}" for k, v in r["bars"].items()))
        wm.cell(row=i, column=6, value=r["bars_total"])
        wm.cell(row=i, column=7, value=r["bar_m"]).number_format = "#,##0.00"
        wm.cell(row=i, column=8, value=f"=IF(G{i}=0,0,D{i}/G{i})").number_format = "0.0%"
        wm.cell(row=i, column=9, value=r["reusable_offcuts"])
        wm.cell(row=i, column=10, value=r["kg"]).number_format = "#,##0.0"
        wm.cell(row=i, column=11, value=f"={price(f'A{i}')}").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
        wm.cell(row=i, column=12, value=f"=G{i}*K{i}").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
        wm.cell(row=i, column=13, value=r["lower_bound_bars"])
    nm = len(q["consolidated"]) + 1
    t = nm + 1
    wm.cell(row=t, column=1, value="TOTAL").font = Font(name=F, bold=True)
    for col, fmt in ((3, "#,##0"), (4, "#,##0.00"), (6, "#,##0"), (7, "#,##0.00"), (10, "#,##0.0"),
                     (12, 'R$ #,##0.00;(R$ #,##0.00);"-"'), (13, "#,##0")):
        L = get_column_letter(col)
        c = wm.cell(row=t, column=col, value=f"=SUM({L}2:{L}{nm})")
        c.number_format = fmt
        c.font = Font(name=F, bold=True)
    wm.cell(row=t, column=8, value=f"=IF(G{t}=0,0,D{t}/G{t})").number_format = "0.0%"
    _widths(wm, [20, 12, 8, 12, 26, 10, 12, 13, 13, 11, 12, 14, 12])

    # ------------------------------------------------ Placas
    ws_ = wb.create_sheet("Placas")
    _hdr(ws_, 1, ["Item", "Posições", "Área bruta (m²)", "Área líquida (m²)", "Placas a comprar", "Preço (R$/placa)",
                  "Custo (R$)"])
    for i, r in enumerate(q["sheets"], start=2):
        ws_.cell(row=i, column=1, value=r["item"])
        ws_.cell(row=i, column=2, value=r["positions"])
        ws_.cell(row=i, column=3, value=r["area_m2"]).number_format = "#,##0.00"
        ws_.cell(row=i, column=4, value=r["net_m2"]).number_format = "#,##0.00"
        ws_.cell(row=i, column=5, value=r["sheets_to_buy"])
        ws_.cell(row=i, column=6, value=f"={price(f'A{i}')}").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
        ws_.cell(row=i, column=7, value=f"=E{i}*F{i}").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
    ns = len(q["sheets"]) + 1
    ws_.cell(row=ns + 1, column=1, value="TOTAL").font = Font(name=F, bold=True)
    ws_.cell(row=ns + 1, column=5, value=f"=SUM(E2:E{ns})")
    ws_.cell(row=ns + 1, column=7, value=f"=SUM(G2:G{ns})").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
    ws_.cell(row=ns + 3, column=1, value="Placas a comprar = área bruta ÷ área da placa × 1,10 (perda de corte). Hipótese a validar com o plano de corte 2D da fábrica.")
    _widths(ws_, [26, 10, 14, 14, 14, 14, 14])

    # ------------------------------------------------ Ferragens
    wf = wb.create_sheet("Ferragens")
    _hdr(wf, 1, ["Item", "Unidade", "Quantidade calculada", "Perda", "Quantidade a comprar", "Preço (R$/un)", "Custo (R$)"])
    for i, r in enumerate(q["hardware"], start=2):
        wf.cell(row=i, column=1, value=r["item"])
        wf.cell(row=i, column=2, value=r["unit"])
        wf.cell(row=i, column=3, value=r["qty"])
        wf.cell(row=i, column=4, value=r["waste"]).number_format = "0%"
        wf.cell(row=i, column=5, value=f"=ROUNDUP(C{i}*(1+D{i}),0)")
        wf.cell(row=i, column=6, value=f"={price(f'A{i}')}").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
        wf.cell(row=i, column=7, value=f"=E{i}*F{i}").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
    nf = len(q["hardware"]) + 1
    wf.cell(row=nf + 1, column=1, value="TOTAL").font = Font(name=F, bold=True)
    wf.cell(row=nf + 1, column=7, value=f"=SUM(G2:G{nf})").number_format = 'R$ #,##0.00;(R$ #,##0.00);"-"'
    _widths(wf, [26, 9, 18, 8, 18, 14, 14])

    # ------------------------------------------------ Pecas
    wpc = wb.create_sheet("Pecas")
    cols = ["ID", "Função", "Item", "Grupo", "Pai (painel/treliça)", "Corte (mm)", "Camadas", "Ângulo A", "Ângulo B",
            "x0", "x1", "z0", "z1", "Especial", "Regra", "Observação", "Peso (kg)"]
    _hdr(wpc, 1, cols)
    for i, m in enumerate(res.members, start=2):
        vals = [m.id, m.role, m.item, m.group, m.parent, m.cut_length, m.plies, m.angle_a, m.angle_b, m.x0, m.x1,
                m.z0, m.z1, "sim" if m.special else "", m.rule, m.note, m.weight_kg]
        for j, v in enumerate(vals, start=1):
            wpc.cell(row=i, column=j, value=v)
    wpc.freeze_panes = "A2"
    wpc.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(res.members) + 1}"
    _widths(wpc, [40, 14, 18, 8, 28, 10, 8, 8, 8, 9, 9, 9, 9, 8, 24, 36, 9])

    # ------------------------------------------------ Plano de corte
    wc = wb.create_sheet("Plano_de_Corte")
    _hdr(wc, 1, ["Item", "Barra nº", "Comprimento da barra", "Cortes (mm)", "Usado (mm)", "Sobra (mm)", "IDs das peças"])
    r = 2
    kerf = None
    for item, plan in q["cut_plan"].items():
        for bi, b in enumerate(plan["bars"], start=1):
            used = sum(c[1] for c in b.cuts)
            wc.cell(row=r, column=1, value=item)
            wc.cell(row=r, column=2, value=bi)
            wc.cell(row=r, column=3, value=b.stock)
            wc.cell(row=r, column=4, value=" + ".join(f"{c[1]:.0f}" for c in b.cuts))
            wc.cell(row=r, column=5, value=round(used, 1))
            wc.cell(row=r, column=6, value=f"=C{r}-E{r}")
            wc.cell(row=r, column=7, value=", ".join(c[0].split("-", 2)[-1] for c in b.cuts))
            r += 1
    wc.cell(row=r + 1, column=1, value="Sobra = barra − soma dos cortes (a perda de serra é descontada na alocação).")
    _widths(wc, [20, 8, 12, 40, 11, 11, 90])

    # ------------------------------------------------ Paineis
    wpn = wb.create_sheet("Paineis")
    _hdr(wpn, 1, ["Painel", "Parede", "Nível", "Tipo", "Comprimento", "Altura", "Peso (kg)", "Peças", "Placas",
                  "Assinatura", "Especial"])
    for i, p in enumerate(q["panels"], start=2):
        for j, v in enumerate([p["id"], p["wall"], p["level"], p["type"], p["length"], p["height"], p["weight_kg"],
                               p["members"], p["sheets"], p["signature"], "sim" if p["special"] else ""], start=1):
            wpn.cell(row=i, column=j, value=v)
    _widths(wpn, [28, 8, 6, 7, 12, 8, 10, 7, 7, 14, 8])

    # ------------------------------------------------ Validacao
    wv = wb.create_sheet("Validacao")
    _hdr(wv, 1, ["Código", "Severidade", "Elemento", "Mensagem"])
    for i, it in enumerate(sorted(res.issues, key=lambda x: ({"error": 0, "warning": 1, "info": 2}[x.severity], x.code)), start=2):
        for j, v in enumerate([it.code, it.severity, it.element, it.message], start=1):
            c = wv.cell(row=i, column=j, value=v)
            if it.severity == "error":
                c.font = Font(name=F, color="C62828")
    _widths(wv, [8, 10, 34, 120])

    # ------------------------------------------------ Resumo (primeira aba)
    wr = wb.create_sheet("Resumo", 0)
    wr["A1"] = f"{res.project} · {res.project_name}"
    wr["A1"].font = Font(name=F, bold=True, size=14)
    wr["A2"] = f"Sistema {res.system} · ruleset {res.ruleset} · execução {res.run_id}"
    wr["A3"] = "STATUS: " + ("APROVADO" if res.ruleset_approved else "RASCUNHO — regras e catálogo sem aprovação do engenheiro; não usar para compra ou fábrica")
    wr["A3"].font = Font(name=F, bold=True, color="2E7D32" if res.ruleset_approved else "C62828")
    lines = [
        ("Painéis", res.stats["panels"], None),
        ("Painéis diferentes", res.stats["unique_panels"], None),
        ("Peças", res.stats["members"], None),
        ("Ligações", res.stats["connections"], None),
        ("Treliças", res.stats["trusses"], None),
        ("Metros de corte (lineares)", f"=Materiais!D{nm + 1}", "#,##0.0"),
        ("Barras a comprar", f"=Materiais!F{nm + 1}", "#,##0"),
        ("Aproveitamento das barras", f"=Materiais!H{nm + 1}", "0.0%"),
        ("Placas a comprar", f"=Placas!E{ns + 1}", "#,##0"),
        ("Peso das peças lineares (kg)", f"=Materiais!J{nm + 1}", "#,##0"),
        ("Custo lineares (R$)", f"=Materiais!L{nm + 1}", 'R$ #,##0.00;(R$ #,##0.00);"-"'),
        ("Custo placas (R$)", f"=Placas!G{ns + 1}", 'R$ #,##0.00;(R$ #,##0.00);"-"'),
        ("Custo ferragens (R$)", f"=Ferragens!G{nf + 1}", 'R$ #,##0.00;(R$ #,##0.00);"-"'),
        ("CUSTO TOTAL DE MATERIAL (R$)", "=B15+B16+B17", 'R$ #,##0.00;(R$ #,##0.00);"-"'),
        ("Erros de validação", res.stats["issues"].get("error", 0), None),
        ("Avisos", res.stats["issues"].get("warning", 0), None),
    ]
    for i, (k, v, fmt) in enumerate(lines, start=5):
        wr.cell(row=i, column=1, value=k)
        c = wr.cell(row=i, column=2, value=v)
        if fmt:
            c.number_format = fmt
        if isinstance(v, str) and v.startswith("=") and "!" in v:
            c.font = Font(name=F, color="008000")
    wr["A18"].font = Font(name=F, bold=True)
    wr["B18"].font = Font(name=F, bold=True)
    wr["A23"] = "Como usar: preencha os preços na aba Precos (células amarelas). Os custos se atualizam sozinhos."
    wr["A24"] = "Texto verde = vínculo com outra aba. Texto azul = entrada manual."
    _widths(wr, [34, 22])
    for ws in wb.worksheets:
        _font_all(ws)
    wb.save(path)
