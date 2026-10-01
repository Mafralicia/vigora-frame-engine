"""Atualização por diff (seção 13) e saídas de fábrica."""
import json
import subprocess
import sys

import ezdxf
from openpyxl import load_workbook

from vigora_frame.builder import load_project, run
from vigora_frame.cli import export_all
from vigora_frame.diff import diff
from conftest import ROOT


def test_moving_one_window_changes_only_its_panel(examples_dir):
    p1 = load_project(examples_dir / "casa_terrea_wood.json")
    p2 = p1.model_copy(deep=True)
    j2 = next(o for o in p2.walls[0].openings if o.id == "J2")
    j2.offset += 200
    a, b = run(p1), run(p2)
    d = diff(a, b)
    assert d["panels_changed"] == ["CT01-L1-W01-PN02"]
    touched = {m for m in d["added"] + d["removed"] + [c[0] for c in d["changed"]]}
    assert touched and all("-W01-PN02-" in m for m in touched)
    assert d["unchanged"] > 500


def test_export_all(tmp_path, examples_dir):
    s = export_all(str(examples_dir / "casa_terrea_wood.json"), str(tmp_path))
    for f in ("resultado.json", "quantitativo.xlsx", "pecas.csv", "lista_corte.csv", "maquina_generico.csv",
              "paineis_e_trelicas.dxf", "etiquetas_paineis.pdf", "etiquetas_pecas.zpl", "pranchas_A1.pdf", "pranchas_rapidas_A3.pdf", "resumo.json"):
        assert (tmp_path / f).stat().st_size > 0, f
    wb = load_workbook(tmp_path / "quantitativo.xlsx")
    assert wb.sheetnames == ["Resumo", "Precos", "Materiais", "Placas", "Ferragens", "Pecas", "Plano_de_Corte",
                             "Paineis", "Validacao"]
    doc = ezdxf.readfile(tmp_path / "paineis_e_trelicas.dxf")
    zpl = (tmp_path / "etiquetas_pecas.zpl").read_text(encoding="utf-8")
    r = json.loads((tmp_path / "resultado.json").read_text(encoding="utf-8"))
    pids = {p["id"] for p in r["panels"]}
    expected = sum(1 for m in r["members"] if m["parent"] in pids) + sum(1 for s_ in r["sheets"] if s_["parent"] in pids) \
        + len(r["panels"]) + sum(len(t["member_ids"]) + 1 for t in r["trusses"])
    assert len(doc.modelspace()) == expected
    n = sum(1 for m in r["members"] if m["frame"] in ("panel", "truss"))
    assert zpl.count("^XA") == n
    assert s["pages"] >= 8


def test_xlsx_recalculates_without_errors(tmp_path, examples_dir):
    export_all(str(examples_dir / "pavilhao_curvo.json"), str(tmp_path), drawings=False)
    out = subprocess.run([sys.executable, "/mnt/skills/public/xlsx/scripts/recalc.py",
                          str(tmp_path / "quantitativo.xlsx"), "90"], capture_output=True, text=True)
    res = json.loads(out.stdout)
    assert res["status"] == "success" and res["total_errors"] == 0
